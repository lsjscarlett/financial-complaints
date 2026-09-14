"""Build blinded rating sheets for the human rating pass (see paper/human_rating_protocol.md).

Writes analysis/human_rating/{rater_A.xlsx, rater_B.xlsx, calibration.xlsx, key.csv,
calibration_key.csv, rubric.md}. Raters get item_id, the complaint as the models saw it, and the
reply; the key maps item_id -> (Row, Model, cell). Order is randomised per rater.

Environment: HR_COMPLAINTS (default 80), HR_SEED (default 7).
"""

import hashlib
import os
import random

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, "dataset")
OUT = os.path.join(HERE, "human_rating")
os.makedirs(OUT, exist_ok=True)

N = int(os.getenv("HR_COMPLAINTS") or "80")
SEED = int(os.getenv("HR_SEED") or "7")
CELLS = {"v1_terse": "V1", "v2_empathetic": "V2 / A00", "v3_structured": "V3", "f_000_base": "000",
         "f_A0C_empathetic_constraints": "A0C", "f_ABC_empathetic_format_constraints": "ABC"}
MODELS = ["ChatGPT", "Mistral"]
SCORE_COLS = ["acknowledgement", "concreteness", "tone", "grounding", "overall"]
FLAG_COLS = ["has_placeholder", "promises_outcome", "admits_liability", "would_send"]
NARRATIVE_CHARS = 1500


def judge_sample_order():
    """Reproduce the stratified order used by llm_judge.py so the first N complaints are the
    prefix of the 300-complaint judge sample."""
    comp = pd.read_csv(os.path.join(DATA, "complaints_10k.csv"), dtype=str, keep_default_na=False)
    comp = comp.rename(columns={"row_id": "Row"})

    def family(p):
        p = p.lower()
        for key, name in (("credit report", "Credit reporting"), ("credit card", "Credit card / prepaid"),
                          ("prepaid", "Credit card / prepaid"), ("payday", "Payday / personal loan"),
                          ("personal loan", "Payday / personal loan"), ("checking", "Checking / savings")):
            if key in p:
                return name
        return p
    comp["family"] = comp["Product"].map(family)
    rng = random.Random(42)
    picked = []
    shares = comp["family"].value_counts(normalize=True)
    for fam, share in shares.items():
        rows = sorted(comp[comp["family"] == fam]["Row"])
        rng.shuffle(rows)
        picked += rows[: max(1, round(share * 300))]
    picked = picked[:300]
    # interleave families so a prefix stays stratified: sort by within-family rank
    ranks = {}
    fam_of = comp.set_index("Row")["family"]
    counters = {}
    for r in picked:
        f = fam_of[r]
        counters[f] = counters.get(f, 0) + 1
        ranks[r] = counters[f] / shares[f]
    picked = sorted(picked, key=lambda r: ranks[r])
    return picked, comp.set_index("Row")


def build_complaint(row):
    issue, sub = row["Issue"], row["Sub-issue"]
    if sub.lower() == issue.lower():
        sub = ""
    narrative = row["Consumer complaint narrative"]
    if len(narrative) > NARRATIVE_CHARS:
        cut = narrative[:NARRATIVE_CHARS]
        narrative = cut[: cut.rfind(" ")] + " [...]"
    lines = [f"Issue: {issue}"]
    if sub:
        lines.append(f"Sub-issue: {sub}")
    lines.append(f"Complaint: {narrative}")
    return "\n".join(lines)


def load_replies():
    frames = []
    for name in ("llm_responses_long.csv", "llm_responses_study2_long.csv"):
        p = os.path.join(DATA, name)
        if not os.path.exists(p):
            p += ".gz"
        frames.append(pd.read_csv(p, dtype=str, keep_default_na=False))
    d = pd.concat(frames, ignore_index=True)
    return d[(d["Is_Error"].str.lower() != "true") & d["Model"].isin(MODELS) & d["Prompt_Variant"].isin(CELLS)]


def item_id(row, model, variant):
    return "R" + hashlib.sha1(f"{SEED}:{row}:{model}:{variant}".encode()).hexdigest()[:8]


def make_sheet(items, path, rater_seed):
    rng = random.Random(rater_seed)
    order = list(range(len(items)))
    rng.shuffle(order)
    rows = []
    for k, i in enumerate(order, 1):
        it = items[i]
        rows.append({"seq": k, "item_id": it["item_id"], "complaint_id": it["Row"],
                     "complaint": it["complaint"], "reply": it["reply"],
                     **{c: "" for c in SCORE_COLS}, **{c: "" for c in FLAG_COLS}, "note": ""})
    df = pd.DataFrame(rows)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, index=False, sheet_name="ratings")
        ws = xw.sheets["ratings"]
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.styles import Alignment
        from openpyxl.utils import get_column_letter
        n = len(df) + 1
        cols = list(df.columns)
        dv_score = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
        dv_flag = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
        ws.add_data_validation(dv_score); ws.add_data_validation(dv_flag)
        for c in SCORE_COLS:
            L = get_column_letter(cols.index(c) + 1); dv_score.add(f"{L}2:{L}{n}")
        for c in FLAG_COLS:
            L = get_column_letter(cols.index(c) + 1); dv_flag.add(f"{L}2:{L}{n}")
        widths = {"seq": 6, "item_id": 12, "complaint_id": 12, "complaint": 70, "reply": 70, "note": 30}
        for c in cols:
            ws.column_dimensions[get_column_letter(cols.index(c) + 1)].width = widths.get(c, 14)
        for r in range(2, n + 1):
            for c in ("complaint", "reply"):
                ws.cell(row=r, column=cols.index(c) + 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "F2"
    return df


def main():
    order, comp = judge_sample_order()
    main_rows = order[:N]
    calib_rows = order[N:N + 5]
    replies = load_replies().set_index(["Row", "Model", "Prompt_Variant"])["Response"]

    def items_for(rows, per_cell_models):
        out = []
        for r in rows:
            complaint = build_complaint(comp.loc[r])
            for v in CELLS:
                for m in per_cell_models(v):
                    try:
                        text = replies.loc[(r, m, v)]
                    except KeyError:
                        continue
                    out.append({"item_id": item_id(r, m, v), "Row": r, "Model": m, "Prompt_Variant": v,
                                "cell": CELLS[v], "complaint": complaint, "reply": text})
        return out

    items = items_for(main_rows, lambda v: MODELS)
    rng = random.Random(SEED)
    calib = items_for(calib_rows, lambda v: [rng.choice(MODELS)])

    pd.DataFrame(items)[["item_id", "Row", "Model", "Prompt_Variant", "cell"]].to_csv(os.path.join(OUT, "key.csv"), index=False)
    pd.DataFrame(calib)[["item_id", "Row", "Model", "Prompt_Variant", "cell"]].to_csv(os.path.join(OUT, "calibration_key.csv"), index=False)
    make_sheet(items, os.path.join(OUT, "rater_A.xlsx"), SEED * 10 + 1)
    make_sheet(items, os.path.join(OUT, "rater_B.xlsx"), SEED * 10 + 2)
    make_sheet(calib, os.path.join(OUT, "calibration.xlsx"), SEED * 10 + 3)

    # rubric for the raters: the anchors section of the protocol
    proto = open(os.path.join(ROOT, "paper", "human_rating_protocol.md")).read()
    start = proto.index("## 3. The rubric"); end = proto.index("## 4. Procedure")
    open(os.path.join(OUT, "rubric.md"), "w").write("# Rating rubric\n\n" + proto[start + len("## 3. The rubric"):end].strip() + "\n")

    print(f"{len(main_rows)} complaints, {len(items)} main items, {len(calib)} calibration items")
    print(pd.DataFrame(items).groupby(["Model", "cell"]).size().unstack(0).to_string())
    print("wrote", OUT)


if __name__ == "__main__":
    main()
